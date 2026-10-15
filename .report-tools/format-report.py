# เครื่องมือจัดรูปแบบและตรวจเอกสารรายงาน
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
out = root / 'docs/เอกสารรายงาน/รายงานจากระบบจริง'
file = out / 'รายงานระบบจัดการโรงแรม-Hotel-Lobby.docx'
doc = Document(file)
doc.core_properties.author = 'นายศุภวัฒน์ ข่ายทอง'
doc.core_properties.title = 'ระบบจัดการโรงแรม Hotel-Lobby'
doc.core_properties.subject = 'รายงานโครงงาน 5 บท'
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'TH SarabunPSK'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(16)
    cs_size = OxmlElement('w:szCs'); cs_size.set(qn('w:val'), '32')
    style.element.get_or_add_rPr().append(cs_size)
    borders = style.element.xpath('./w:pPr/w:pBdr')
    for border in borders:
        border.getparent().remove(border)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('ascii','hAnsi','cs','eastAsia'):
        fonts.set(qn('w:'+attr), 'TH SarabunPSK')

for p in doc.paragraphs:
    plain = p.text.replace('\u200b', '')
    if plain == 'คณะผู้จัดทำ':
        p.runs[0].text = 'ผู้จัดทำ'
    p.paragraph_format.widow_control = True
    if not p._p.xpath('.//w:drawing'):
        p.paragraph_format.line_spacing = 1.15
    if plain.startswith('ระบบจัดการโรงแรม Hotel-Lobby'):
        p.style = doc.styles['Title']
    if plain.startswith('บทที่ '):
        p.paragraph_format.keep_with_next = True
    if p.runs and p.runs[0].bold and p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        p.style = doc.styles['Heading 2']
        p.paragraph_format.keep_with_next = True
    if p.paragraph_format.first_line_indent and p.paragraph_format.first_line_indent > 0:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.color.rgb = RGBColor(0,0,0)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'th-TH')
        lang.set(qn('w:eastAsia'), 'th-TH')
        lang.set(qn('w:bidi'), 'th-TH')
        run._r.get_or_add_rPr().append(lang)

# Bind chapter titles and reference headings to navigation and TOC fields.
titles = {'บทนำ','แนวคิดและเทคโนโลยีที่เกี่ยวข้อง','การออกแบบและพัฒนาระบบ',
          'ผลการพัฒนาและหลักฐานการทดสอบ','สรุปและข้อเสนอแนะ','บรรณานุกรม'}
for p in doc.paragraphs:
    plain = p.text.replace('\u200b', '')
    if plain in titles:
        p.style = doc.styles['Heading 1']
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(20)
        previous = p._p.getprevious()
        if previous is not None and previous.tag == qn('w:p'):
            oldtext = ''.join(previous.itertext()).replace('\u200b','')
            # Read only text runs, avoiding python-docx's duplicated itertext representation.
            oldtext = ''.join(previous.xpath('.//w:t/text()')).replace('\u200b','')
            if oldtext.startswith('บทที่ '):
                number = oldtext.split()[-1]
                p.text = f'บทที่ {number}\n{plain}'
                for run in p.runs:
                    run.font.name = 'TH SarabunPSK'; run.font.size = Pt(20); run.bold = True
                    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                    for attr in ('ascii','hAnsi','cs','eastAsia'):
                        fonts.set(qn('w:'+attr), 'TH SarabunPSK')
                previous.getparent().remove(previous)

# Keep the concluding chapter together without a continuation page of a few lines.
in_conclusion = False
for p in doc.paragraphs:
    plain = p.text.replace('\u200b','')
    if plain.startswith('บทที่ 5'):
        in_conclusion = True
    if plain == 'บรรณานุกรม':
        in_conclusion = False
    if in_conclusion and p.paragraph_format.first_line_indent:
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)

# Keep a compact one-page TOC, updated and paginated in Word.
toc_heading = next(p for p in doc.paragraphs if p.text.replace('\u200b','') == 'สารบัญ')
node = toc_heading._p.getnext()
while node is not None and node.tag == qn('w:p') and not node.xpath('./w:pPr/w:sectPr'):
    next_node = node.getnext()
    node.getparent().remove(node)
    node = next_node
toc = OxmlElement('w:p')
toc_heading._p.addnext(toc)
for label, page in reversed([('คำนำ','ก'), ('สารบัญ','ข')]):
    row = OxmlElement('w:p')
    props = OxmlElement('w:pPr'); tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'),'right'); tab.set(qn('w:leader'),'dot'); tab.set(qn('w:pos'),'8662')
    tabs.append(tab); props.append(tabs); row.append(props)
    for value in (label, '\t', page):
        run = OxmlElement('w:r')
        if value == '\t': run.append(OxmlElement('w:tab'))
        else:
            value_node = OxmlElement('w:t'); value_node.text = value; run.append(value_node)
        row.append(run)
    toc_heading._p.addnext(row)
begin = OxmlElement('w:r'); char = OxmlElement('w:fldChar'); char.set(qn('w:fldCharType'),'begin'); begin.append(char); toc.append(begin)
ins = OxmlElement('w:r'); instruction = OxmlElement('w:instrText'); instruction.set(qn('xml:space'),'preserve'); instruction.text = ' TOC \\o "1-1" \\h \\z \\u '; ins.append(instruction); toc.append(ins)
sep = OxmlElement('w:r'); char = OxmlElement('w:fldChar'); char.set(qn('w:fldCharType'),'separate'); sep.append(char); toc.append(sep)
end = OxmlElement('w:r'); char = OxmlElement('w:fldChar'); char.set(qn('w:fldCharType'),'end'); end.append(char); toc.append(end)
for s in doc.sections:
    s.left_margin = Cm(3.18); s.right_margin = Cm(2.54)
    s.top_margin = Cm(2.54); s.bottom_margin = Cm(2.54)
# Thai uses complex-script size and weight in Word.
for p in doc.paragraphs:
    for run in p.runs:
        props = run._r.get_or_add_rPr()
        size = props.find(qn('w:sz'))
        if size is not None:
            old = props.find(qn('w:szCs'))
            if old is not None: props.remove(old)
            cs_size = OxmlElement('w:szCs'); cs_size.set(qn('w:val'),size.get(qn('w:val')))
            props.append(cs_size)
        if run.bold:
            bold = OxmlElement('w:bCs'); props.append(bold)
# A standalone break paragraph can create an empty page at a page boundary.
for p in list(doc.paragraphs):
    if not p.text and p._p.xpath('.//w:br[@w:type="page"]'):
        next_node = p._p.getnext()
        if next_node is not None and next_node.tag == qn('w:p'):
            props = next_node.find(qn('w:pPr'))
            if props is None:
                props = OxmlElement('w:pPr'); next_node.insert(0,props)
            props.append(OxmlElement('w:pageBreakBefore'))
            p._p.getparent().remove(p._p)
doc.save(file)
print('Formatted', file)
